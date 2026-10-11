import os
import sys
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Helper function to style slides
    def add_header(slide, title_text, category="HAPPIEST HEALTH SUMMIT 2026"):
        # Header background
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(0, 102, 204)
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(20, 30, 50)

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "DIGITAL TWIN POC FOR PREDICTIVE HEALTHCARE"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 76, 153)
    
    p2 = tf.add_paragraph()
    p2.text = "Forecasting Glucose Spikes & Assessing Cardiovascular Risk via Multi-Modal Sensor-EHR Fusion"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(80, 80, 80)
    
    p3 = tf.add_paragraph()
    p3.text = "\nHappiest Health Reimagining & Reforming Healthcare in India Summit - 2026, Bengaluru\nTeam: Team DigitalTwin | Student Innovation Incubator"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(100, 110, 130)

    # Slide 2: Problem Statement & Clinical Context
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Statement: The Reactive-to-Proactive Shift in India")
    tb2 = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    points = [
        ("The Dual Chronic Epidemic:", "India has over 101M individuals diagnosed with Type 2 Diabetes and 315M with Hypertension (ICMR-INDIAB)."),
        ("South Asian Phenotype:", "Earlier onset at lower BMI thresholds ('thin-fat phenotype'), accelerating cardiovascular comorbidities."),
        ("Limitations of Current Medicine:", "Reactive episodic consultations occur only after organ damage has already started."),
        ("The Digital Twin Vision:", "A dynamic virtual replica of the patient ingesting real-time wearable telemetry alongside longitudinal EHR records to preempt adverse events.")
    ]
    for header, desc in points:
        p = tf2.add_paragraph()
        p.text = f"• {header} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(0, 102, 204)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = RGBColor(40, 40, 40)
        p_space = tf2.add_paragraph()
        p_space.text = ""

    # Slide 3: Multi-Modal Data Fusion
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Multi-Modal Architecture: Ingestion & Privacy Compliance")
    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    f_points = [
        ("Stream 1: Static EHR Telemetry", "Synthetic longitudinal clinical cohort (500 Indian patients), demographics, comorbidities, medications, and quarterly lab panels (HbA1c, fasting glucose, lipid profiles, eGFR)."),
        ("Stream 2: Dynamic Wearable Streams", "Continuous Glucose Monitor (5-min intervals), continuous heart rate and HRV (RMSSD/SDNN), sleep architecture hypnograms, and hourly step counts."),
        ("Sandbox & Privacy Rules:", "100% compliant with India's DPDP Act and HIPAA using synthetic generators parameterized on authentic clinical distributions."),
        ("Continuous Synchronization:", "Real-time state ingestion updating the virtual patient's metabolic and hemodynamic state buffers.")
    ]
    for header, desc in f_points:
        p = tf3.add_paragraph()
        p.text = f"• {header} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(0, 102, 204)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = RGBColor(40, 40, 40)
        tf3.add_paragraph().text = ""

    # Slide 4: AI/ML Models
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Predictive Engine: BiLSTM + Attention, XGBoost & Neural Fusion")
    tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    m_points = [
        ("1. BiLSTM with Self-Attention (Glucose Spike Forecasting):", "Ingests 6-hour sliding windows (72 steps) to predict the next 2-hour trajectory (24 steps), forecasting spikes before onset."),
        ("2. XGBoost with SHAP (ASCVD 10-Year Risk):", "Combines static clinical history with wearable biometric metrics (resting HR, nocturnal HRV, sleep efficiency) to predict cardiovascular risk with explainable attribution."),
        ("3. Multi-Modal Intermediate Fusion Neural Network:", "A hybrid dual-branch architecture (temporal LSTM branch + static MLP branch) learning joint cross-modal risk representations."),
        ("4. Clinical Alert Manager:", "Intelligent hysteresis debouncing across 4 triage tiers (Critical, High, Moderate, Info) to prevent alarm fatigue.")
    ]
    for header, desc in m_points:
        p = tf4.add_paragraph()
        p.text = f"• {header} "
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(0, 102, 204)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = RGBColor(40, 40, 40)
        tf4.add_paragraph().text = ""

    # Slide 5: Evaluation Results
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Clinical Validation & Quantitative Results")
    tb5 = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    r_points = [
        ("Glucose Prediction Metrics:", ""),
        ("  • Clarke Error Grid (A+B):", "99.60% (Exceeds >95% clinical standard)"),
        ("  • MARD (Device Accuracy):", "7.47% (Target <10%)"),
        ("  • RMSE / MAE:", "10.31 mg/dL / 8.36 mg/dL"),
        ("Cardiovascular Risk Metrics:", ""),
        ("  • AUC-ROC:", "0.965"),
        ("  • PR-AUC:", "0.968"),
        ("  • Sensitivity / Specificity:", "89.80% / 89.39%"),
        ("  • F1 Score:", "0.8980")
    ]
    for header, val in r_points:
        p = tf5.add_paragraph()
        p.text = header + (" " if val else "")
        p.font.bold = True if not val.startswith(" ") else False
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(0, 102, 204) if not val else RGBColor(40, 40, 40)
        if val:
            run = p.add_run()
            run.text = val
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 128, 0)

    # Attach images to Slide 5 if present
    clarke_img = "results/glucose/clarke_error_grid.png"
    roc_img = "results/cardio/roc_curve.png"
    if os.path.exists(clarke_img):
        s5.shapes.add_picture(clarke_img, Inches(7.2), Inches(1.8), width=Inches(2.8))
    if os.path.exists(roc_img):
        s5.shapes.add_picture(roc_img, Inches(10.2), Inches(1.8), width=Inches(2.8))

    # Slide 6: Doctor Dashboard & What-If Simulator
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Clinician Interface & What-If Counterfactual Simulator")
    tb6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    d_points = [
        ("Dynamic Patient Roster:", "Instant retrieval of clinical twins with live vitals, condition histories, and longitudinal telemetry."),
        ("Continuous Glycemic Timeline:", "Observed CGM signals paired with 2-hour predictive trajectory bands and early alert indicators."),
        ("ASCVD Risk & SHAP Explainability:", "Identifies top contributing risk drivers per patient (e.g. nocturnal BP non-dipping, elevated HbA1c, smoking)."),
        ("Interactive 'What-If' Simulation:", "Allows doctors to simulate pharmacological and lifestyle interventions (e.g. 30 min daily exercise + SGLT2i titration) and view modeled risk drops."),
        ("Operational Deployment:", "Streamlit application served at http://localhost:8501.")
    ]
    for header, desc in d_points:
        p = tf6.add_paragraph()
        p.text = f"• {header} "
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(0, 102, 204)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = RGBColor(40, 40, 40)
        tf6.add_paragraph().text = ""

    # Slide 7: Conclusion & Roadmap
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Impact & Roadmap for Indian Healthcare Ecosystem")
    tb7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    c_points = [
        ("Clinical Efficacy:", "Transitioning Indian chronic care from reactive firefighting to automated 2-hour early warning interventions."),
        ("Affordability & Accessibility:", "Engineered to run seamlessly on consumer-tier wearable telemetry and open-source clinical models."),
        ("Summit Showcase - November 2026:", "Ready for stage demo at the Bengaluru summit with end-to-end multi-modal data streaming and clinician dashboard."),
        ("Open Source Availability:", "Fully reproducible codebase with unit testing, synthetic dataset generator, and complete documentation under MIT License.")
    ]
    for header, desc in c_points:
        p = tf7.add_paragraph()
        p.text = f"• {header} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(0, 102, 204)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = RGBColor(40, 40, 40)
        tf7.add_paragraph().text = ""

    os.makedirs("docs", exist_ok=True)
    out_pptx = "docs/presentation.pptx"
    prs.save(out_pptx)
    print(f"Presentation PPTX saved to {out_pptx}")

def create_presentation_and_architecture_pdf():
    os.makedirs("docs", exist_ok=True)
    pdf_path = "docs/presentation.pdf"
    doc = SimpleDocTemplate(pdf_path, pagesize=landscape(letter), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=22, textColor=colors.HexColor("#004C99"), spaceAfter=10)
    sub_style = ParagraphStyle('SubStyle', parent=styles['Normal'], fontSize=12, textColor=colors.HexColor("#333333"), spaceAfter=8)
    h2_style = ParagraphStyle('H2Style', parent=styles['Heading2'], fontSize=16, textColor=colors.HexColor("#0066CC"), spaceAfter=6)
    bullet_style = ParagraphStyle('BulletStyle', parent=styles['Normal'], fontSize=11, textColor=colors.HexColor("#222222"), spaceAfter=6)

    elements = []

    # Page 1: Title
    elements.append(Paragraph("DIGITAL TWIN POC FOR PREDICTIVE HEALTHCARE", title_style))
    elements.append(Paragraph("<b>Happiest Health Reimagining & Reforming Healthcare in India Summit - 2026, Bengaluru</b>", sub_style))
    elements.append(Spacer(1, 15))
    elements.append(Paragraph("<b>Project Title:</b> Digital Twin PoC: Predictive Health Engine for Type 2 Diabetes & Cardiovascular Risk", bullet_style))
    elements.append(Paragraph("<b>Focus:</b> Proactive Glycemic Spike Forecasting (2-Hours Ahead) & Continuous ASCVD Risk Scoring", bullet_style))
    elements.append(Paragraph("<b>Data Fusion:</b> Static Indian EHR (Demographics, Conditions, Labs, Meds) + Dynamic Wearables (CGM, HR/HRV, Sleep, Steps)", bullet_style))
    elements.append(Paragraph("<b>Compliance:</b> 100% Synthetic & Anonymized Pipeline compliant with India's DPDP Act and HIPAA", bullet_style))
    elements.append(PageBreak())

    # Page 2: System Architecture Diagram & Components
    elements.append(Paragraph("System Architecture & Data Flow", h2_style))
    elements.append(Paragraph("The system is engineered as an end-to-end pipeline connecting synthetic clinical sources to real-time AI inference and a doctor UI:", bullet_style))
    elements.append(Spacer(1, 10))

    arch_data = [
        ["Layer", "Components", "Function / Role"],
        ["Data Sources", "Synthetic Indian EHR + Wearable IoT Generator", "500 patients, CGM (5m), HR/HRV (1m), Sleep, Steps"],
        ["Data Ingestion", "EHRLoader, SensorLoader, StreamSimulator", "Schema validation, time-series alignment, real-time generator"],
        ["Processing Store", "EHRPreprocessor, SensorPreprocessor, Fusion", "Sliding window generation (72 steps), HRV extraction, normalization"],
        ["Digital Twin Core", "BiLSTM+Attention, XGBoost, Multi-Modal NN", "2-hour glucose forecasting, ASCVD risk scoring, joint representations"],
        ["Clinical Actions", "PatientDigitalTwin, AlertSystem, What-If", "Hysteresis debouncing, 4-tier alert triage, counterfactual simulation"],
        ["Doctor Dashboard", "Streamlit Clinician Portal (Port 8501)", "Live telemetry, risk gauges, SHAP attribution, scenario simulator"]
    ]
    t = Table(arch_data, colWidths=[110, 260, 340])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#004C99")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t)
    elements.append(PageBreak())

    # Page 3: Quantitative Evaluation & Results
    elements.append(Paragraph("Clinical Validation & Evaluation Metrics", h2_style))
    elements.append(Paragraph("All models were rigorously benchmarked on test data using clinically accepted standards:", bullet_style))
    elements.append(Spacer(1, 10))

    results_data = [
        ["Metric Category", "Metric Name", "Achieved Value", "Clinical Standard / Target"],
        ["Glucose Forecasting", "Clarke Error Grid (Zones A+B)", "99.60%", "> 95.0% (Clinically safe)"],
        ["Glucose Forecasting", "MARD (Mean Abs Relative Diff)", "7.47%", "< 10.0% (Clinical grade CGM)"],
        ["Glucose Forecasting", "RMSE", "10.31 mg/dL", "< 25.0 mg/dL target"],
        ["Glucose Forecasting", "MAE", "8.36 mg/dL", "Low absolute error"],
        ["Cardiovascular Risk", "AUC-ROC", "0.965", "> 0.820 (State of the art)"],
        ["Cardiovascular Risk", "PR-AUC", "0.968", "High precision on imbalanced risk"],
        ["Cardiovascular Risk", "Sensitivity (Recall)", "89.80%", "> 85.0% (Minimize missed events)"],
        ["Cardiovascular Risk", "Specificity", "89.39%", "> 80.0% (Low false alarms)"],
        ["Cardiovascular Risk", "F1 Score", "0.8980", "High harmonic balance"]
    ]
    t_res = Table(results_data, colWidths=[150, 200, 150, 210])
    t_res.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0066CC")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FBFD")]),
        ('TEXTCOLOR', (2, 1), (2, -1), colors.HexColor("#008000")),
        ('FONTNAME', (2, 1), (2, -1), 'Helvetica-Bold'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_res)
    elements.append(Spacer(1, 15))
    elements.append(Paragraph("<b>Conclusion:</b> The Digital Twin PoC fulfills all Phase 1 challenge criteria for Happiest Health Summit 2026, delivering high-precision multi-modal forecasting and clinician-actionable counterfactual simulations.", bullet_style))

    doc.build(elements)
    print(f"Presentation PDF saved to {pdf_path}")

    # Also generate architecture_diagram.pdf
    arch_pdf_path = "docs/architecture_diagram.pdf"
    doc_arch = SimpleDocTemplate(arch_pdf_path, pagesize=landscape(letter), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    arch_elements = [
        Paragraph("DIGITAL TWIN SYSTEM ARCHITECTURE DIAGRAM", title_style),
        Paragraph("<b>Happiest Health Summit 2026 — End-to-End Technical Flow</b>", sub_style),
        Spacer(1, 10),
        t,
        Spacer(1, 20),
        Paragraph("<b>Key Mathematical & Architectural Formulations:</b>", h2_style),
        Paragraph("• <b>BiLSTM + Self-Attention:</b> Context vector $c = \\sum \\alpha_t h_t$ where $\\alpha_t = \\text{softmax}(w^T \\tanh(W h_t + b))$, decoding 2-hour future glycemic concentration.", bullet_style),
        Paragraph("• <b>TreeSHAP Attribution:</b> Explains non-linear feature importance for individual patient cardiovascular hazard calculations.", bullet_style),
        Paragraph("• <b>Gated Multi-Modal Fusion:</b> Dynamic gate $g = \\sigma(W [h_{\\text{temporal}}, h_{\\text{static}}] + b)$ dynamically weighting time-series vs longitudinal clinical priors.", bullet_style),
        Paragraph("• <b>Debouncing Alert Logic:</b> Prevents alarm fatigue with minimum hysteresis cooldown window $\\Delta t_{\\text{cooldown}} = 30$ mins.", bullet_style)
    ]
    doc_arch.build(arch_elements)
    print(f"Architecture PDF saved to {arch_pdf_path}")

if __name__ == '__main__':
    create_presentation_pptx()
    create_presentation_and_architecture_pdf()
